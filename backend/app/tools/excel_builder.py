import os
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../storage/outputs"))
os.makedirs(OUTPUT_DIR, exist_ok=True)

def create_styled_excel(
    filename: str,
    sheets_data: list,
    theme_color: str = "1F4E78" # Default Elegant Navy
) -> dict:
    """
    Creates a professionally styled Excel file.
    
    sheets_data structure:
    [
        {
            "name": "Sheet1",
            "title": "2026年第一季度销售数据汇总",
            "headers": ["月份", "产品", "销量", "单价", "销售额", "达成率"],
            "rows": [
                ["1月", "办公协同套件", 1200, 199.0, "=C4*D4", 0.95],
                ...
            ],
            "summary_row": ["总计", "-", "=SUM(C4:C10)", "-", "=SUM(E4:E10)", "-"] # Optional
        }
    ]
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    title_font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color=theme_color, end_color=theme_color, fill_type="solid")
    
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    # Slightly lighter for headers
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    
    data_font = Font(name="Microsoft YaHei", size=10, color="000000")
    zebra_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    summary_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    summary_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1F4E78")
    
    thin_border_side = Side(style="thin", color="D3D3D3")
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    double_bottom = Side(style="double", color="1F4E78")
    summary_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_bottom)

    for sheet_idx, s_data in enumerate(sheets_data):
        sheet_name = s_data.get("name", f"Sheet{sheet_idx+1}")[:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        current_row = 1
        headers = s_data.get("headers", [])
        num_cols = max(len(headers), 1)

        # 1. Sheet Title Banner (if provided)
        title_text = s_data.get("title")
        if title_text:
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=num_cols)
            title_cell = ws.cell(row=current_row, column=1, value=title_text)
            title_cell.font = title_font
            title_cell.fill = title_fill
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[current_row].height = 40
            current_row += 2 # Leave 1 empty row for breathing room

        # 2. Table Headers
        header_row = current_row
        ws.row_dimensions[header_row].height = 26
        for col_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = data_border
        current_row += 1

        # 3. Data Rows
        rows = s_data.get("rows", [])
        for r_idx, row_values in enumerate(rows):
            ws.row_dimensions[current_row].height = 22
            is_zebra = (r_idx % 2 == 1)
            for c_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.font = data_font
                cell.border = data_border
                if is_zebra:
                    cell.fill = zebra_fill
                
                # Intelligent alignment & format detection
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    # Percentage or Currency check
                    if 0 < val <= 1 and any(keyword in str(headers[c_idx-1] if c_idx-1 < len(headers) else "") for keyword in ["率", "%", "比", "rate", "percent"]):
                        cell.number_format = "0.0%"
                    elif isinstance(val, float):
                        cell.number_format = "#,##0.00"
                    else:
                        cell.number_format = "#,##0"
                elif isinstance(val, str) and val.startswith("="):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    # If formula points to currency/number
                    cell.number_format = "#,##0.00"
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        # 4. Summary Row (Optional)
        summary_row = s_data.get("summary_row")
        if summary_row:
            ws.row_dimensions[current_row].height = 25
            for c_idx, val in enumerate(summary_row, 1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.font = summary_font
                cell.fill = summary_fill
                cell.border = summary_border
                if isinstance(val, (int, float)) or (isinstance(val, str) and val.startswith("=")):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0.00"
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            current_row += 1

        # 5. Auto adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Skip merged title row for width calculation
                if title_text and cell.row == 1:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    # Count Chinese characters as 2 length
                    char_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
                    max_len = max(max_len, char_len)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to file
    safe_name = os.path.basename(filename)
    if not safe_name.endswith(".xlsx"):
        safe_name += ".xlsx"
    file_id = f"{uuid.uuid4().hex[:8]}_{safe_name}"
    file_path = os.path.join(OUTPUT_DIR, file_id)
    wb.save(file_path)

    return {
        "file_id": file_id,
        "filename": safe_name,
        "file_path": file_path,
        "file_size": os.path.getsize(file_path),
        "file_type": "excel"
    }
