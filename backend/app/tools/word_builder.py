import os
import uuid
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../storage/outputs"))
os.makedirs(OUTPUT_DIR, exist_ok=True)

def set_cell_background(cell, hex_color: str):
    """Sets background color of a table cell (e.g., '1F4E78')"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in twentieths of a point (dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_styled_word(
    filename: str,
    title: str,
    subtitle: str = "",
    sections: list = None
) -> dict:
    """
    Creates a professionally formatted Word (.docx) document.
    
    sections structure:
    [
        {
            "heading": "一、 项目背景与概述",
            "level": 1, # 1, 2, or 3
            "content": [
                "第一段正文内容...",
                "第二段正文内容..."
            ],
            "bullets": ["要点1", "要点2"],
            "callout": "核心结论或重点备忘：这里是突出展示的提示块。",
            "table": {
                "headers": ["阶段", "任务", "负责人", "周期"],
                "rows": [
                    ["阶段一", "需求分析与方案设计", "张工", "1周"],
                    ["阶段二", "核心Agent调度与开发", "李工", "2周"]
                ]
            }
        }
    ]
    """
    doc = docx.Document()

    # Base margins
    sections_doc = doc.sections
    for s in sections_doc:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # 1. Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(6)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.font.name = "Microsoft YaHei"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78) # Elegant Deep Navy

    # 2. Subtitle (if any)
    if subtitle:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(20)
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(subtitle)
        sub_run.font.name = "Microsoft YaHei"
        sub_run.font.size = Pt(11)
        sub_run.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)

    # Divider line
    divider = doc.add_paragraph()
    divider.paragraph_format.space_after = Pt(16)
    divider_run = divider.add_run("—" * 38)
    divider_run.font.color.rgb = RGBColor(0xDC, 0xDF, 0xE6)
    divider.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 3. Process Sections
    if sections:
        for sec in sections:
            # Heading
            h_text = sec.get("heading")
            h_level = sec.get("level", 1)
            if h_text:
                h_p = doc.add_paragraph()
                h_p.paragraph_format.space_before = Pt(14)
                h_p.paragraph_format.space_after = Pt(6)
                h_run = h_p.add_run(h_text)
                h_run.font.name = "Microsoft YaHei"
                h_run.font.bold = True
                if h_level == 1:
                    h_run.font.size = Pt(15)
                    h_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
                else:
                    h_run.font.size = Pt(13)
                    h_run.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

            # Text content
            content_list = sec.get("content", [])
            if isinstance(content_list, str):
                content_list = [content_list]
            for para_text in content_list:
                if para_text:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.25
                    run = p.add_run(para_text)
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)

            # Bullet points
            bullets = sec.get("bullets", [])
            for b_text in bullets:
                b_p = doc.add_paragraph(style='List Bullet')
                b_p.paragraph_format.space_after = Pt(4)
                b_run = b_p.add_run(b_text)
                b_run.font.name = "Microsoft YaHei"
                b_run.font.size = Pt(10.5)
                b_run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)

            # Callout box
            callout = sec.get("callout")
            if callout:
                c_table = doc.add_table(rows=1, cols=1)
                c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                c_cell = c_table.cell(0, 0)
                set_cell_background(c_cell, "F2F4F8")
                set_cell_margins(c_cell, top=140, bottom=140, left=200, right=200)
                
                c_p = c_cell.paragraphs[0]
                c_p.paragraph_format.space_before = Pt(0)
                c_p.paragraph_format.space_after = Pt(0)
                
                prefix_run = c_p.add_run("💡 核心提示 / 总结：\n")
                prefix_run.font.name = "Microsoft YaHei"
                prefix_run.font.size = Pt(10)
                prefix_run.font.bold = True
                prefix_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
                
                c_run = c_p.add_run(callout)
                c_run.font.name = "Microsoft YaHei"
                c_run.font.size = Pt(10)
                c_run.font.italic = True
                c_run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

                # spacing after table
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(6)

            # Table
            table_data = sec.get("table")
            if table_data and table_data.get("headers"):
                t_headers = table_data.get("headers", [])
                t_rows = table_data.get("rows", [])
                
                doc_table = doc.add_table(rows=len(t_rows) + 1, cols=len(t_headers))
                doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Header Row
                for col_idx, h_name in enumerate(t_headers):
                    cell = doc_table.cell(0, col_idx)
                    set_cell_background(cell, "2F5597")
                    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                    h_p = cell.paragraphs[0]
                    h_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = h_p.add_run(str(h_name))
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(10)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                # Data Rows
                for r_idx, row_vals in enumerate(t_rows):
                    row_num = r_idx + 1
                    is_zebra = (r_idx % 2 == 1)
                    for col_idx, val in enumerate(row_vals):
                        if col_idx < len(t_headers):
                            cell = doc_table.cell(row_num, col_idx)
                            if is_zebra:
                                set_cell_background(cell, "F2F4F8")
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            c_p = cell.paragraphs[0]
                            run = c_p.add_run(str(val))
                            run.font.name = "Microsoft YaHei"
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

                # spacing after table
                t_spacer = doc.add_paragraph()
                t_spacer.paragraph_format.space_after = Pt(10)

    # Save Word document
    safe_name = os.path.basename(filename)
    if not safe_name.endswith(".docx"):
        safe_name += ".docx"
    file_id = f"{uuid.uuid4().hex[:8]}_{safe_name}"
    file_path = os.path.join(OUTPUT_DIR, file_id)
    doc.save(file_path)

    return {
        "file_id": file_id,
        "filename": safe_name,
        "file_path": file_path,
        "file_size": os.path.getsize(file_path),
        "file_type": "word"
    }
